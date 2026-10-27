"""Editable PowerPoint version of Fig. 1 (A-CBFI pipeline).

Text follows pending item 11 (G3/G4 labels, notation of eq. 3).
Run:  python make_fig1_pptx.py   ->  exp_result/manuscript_assets/fig1_pipeline.pptx
"""
import json
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
OUT = os.path.join(ROOT, "exp_result", "manuscript_assets")
FONT, SIZE = "Arial", 11
FILL, LINE, NOTE = RGBColor(0xEA, 0xF0, 0xF8), RGBColor(0x4C, 0x72, 0xB0), RGBColor(0x55, 0x55, 0x55)

# A run is (text, style); style letters: b bold, i italic, s subscript, p superscript.
R = lambda t, s="": (t, s)
CG4BR = [R("C"), R("G4", "s"), R("br", "p")]


def add_runs(par, runs, size=SIZE, color=None):
    for text, style in runs:
        r = par.add_run()
        r.text = text
        f = r.font
        f.name, f.size, f.bold, f.italic = FONT, Pt(size), "b" in style, "i" in style
        if color is not None:
            f.color.rgb = color
        if "s" in style:
            r._r.get_or_add_rPr().set("baseline", "-25000")
        if "p" in style:
            r._r.get_or_add_rPr().set("baseline", "30000")


def text_box(slide, x, y, w, h, lines, size=SIZE, color=None, align=PP_ALIGN.CENTER):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for k, runs in enumerate(lines):
        par = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        par.alignment = align
        add_runs(par, runs, size, color)
    return tb


def box(slide, x, y, w, h, lines):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.08
    shp.fill.solid(); shp.fill.fore_color.rgb = FILL
    shp.line.color.rgb = LINE; shp.line.width = Pt(1.25)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.08)
    for k, runs in enumerate(lines):
        par = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        par.alignment = PP_ALIGN.CENTER
        if k == 1:
            par.space_before = Pt(4)
        add_runs(par, runs, size=SIZE if k < 2 else SIZE - 1, color=RGBColor(0, 0, 0))
    return shp


def arrow(slide, x1, y, x2):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y), Inches(x2), Inches(y))
    c.line.color.rgb = RGBColor(0, 0, 0); c.line.width = Pt(1.25)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)


def main():
    A = json.load(open(os.path.join(OUT, "assets.json"), encoding="utf-8"))
    share = A["main_total"]["single_share_of_solved"]

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10.4), Inches(3.0)
    s = prs.slides.add_slide(prs.slide_layouts[6])

    steps = [
        [[R("Step 0", "b")], [R("SCM over features", "b")],
         [R("domain DAG "), R("G", "i"), R(",")],
         [R("direction and bound constraints,")],
         [R("abducted noise "), R("û", "i")]],
        [[R("Step 1", "b")], [R("Prescriptive single-lever diagnosis", "b")],
         [R("F1("), R("i", "i"), R(") = max"), R("v", "si"), R("∈", "s"), R("V", "si"), R("i", "si"),
          R(" "), R("g", "i"), R("i", "si"), R("("), R("v", "i"), R(")")],
         [R("for every actionable feature "), R("i", "i")]],
        [[R("Step 2", "b")], [R("Does a single action reach the boundary?", "b")],
         [R("→ cheapest single action")],
         [R("+ menu of substitutable single actions (G4-type)")]],
        [[R("Step 3", "b")], [R("Escalation", "b")],
         [R("joint actions needed")],
         [R("(G3-type, ")] + CG4BR + [R(" > 0):")],
         [R("pairs, then triples")],
         [R("(F2 screening or top-"), R("M", "i"), R(")")]],
    ]
    x0, y0, w, h, gap = 0.3, 0.75, 2.2, 1.75, 0.35
    for k, lines in enumerate(steps):
        box(s, x0 + k * (w + gap), y0, w, h, lines)
        if k:
            arrow(s, x0 + k * (w + gap) - gap + 0.03, y0 + h / 2, x0 + k * (w + gap) - 0.03)

    x2 = x0 + 2 * (w + gap)
    text_box(s, x2 + w - 0.05, y0 + h / 2 - 0.42, gap + 0.1, 0.3, [[R("no")]], size=9, color=NOTE)
    text_box(s, x2 - 0.2, y0 + h + 0.05, w + 0.4, 0.3,
             [[R(f"yes: stop ({100 * share:.0f}% of solved queries)")]], size=9, color=NOTE)
    text_box(s, 0.3, 0.2, 9.8, 0.35,
             [[R("Output: direct actions (", "i"), R("L", "i"), R("active", "si"),
               R(") separated from SCM-propagated downstream changes", "i")]], size=SIZE)

    path = os.path.join(OUT, "fig1_pipeline.pptx")
    prs.save(path)
    print(path)


if __name__ == "__main__":
    main()
